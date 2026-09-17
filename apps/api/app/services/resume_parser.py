import io
import re
from typing import Tuple
import docx
from pypdf import PdfReader

from app.core.exceptions import ValidationError

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


class DocumentParser:
    """Multi-format deterministic text extraction for PDF, DOCX, and TXT resumes."""

    @staticmethod
    def validate_file(file_name: str, file_bytes: bytes) -> str:
        """Validate file size and extension."""
        if len(file_bytes) == 0:
            raise ValidationError("Uploaded file is empty (0 bytes).")

        if len(file_bytes) > MAX_FILE_SIZE_BYTES:
            raise ValidationError(
                f"File size ({len(file_bytes) / (1024 * 1024):.1f}MB) exceeds the maximum allowed 10MB limit."
            )

        lower_name = file_name.lower()
        if lower_name.endswith(".pdf"):
            return "pdf"
        elif lower_name.endswith(".docx"):
            return "docx"
        elif lower_name.endswith(".txt"):
            return "txt"
        else:
            raise ValidationError(
                f"Unsupported file format for '{file_name}'. Only PDF (.pdf), DOCX (.docx), and Plain Text (.txt) are supported."
            )

    @staticmethod
    def parse_pdf(file_bytes: bytes) -> str:
        """Extract text layers from PDF bytes using PyPDF."""
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            text_chunks = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)
            
            combined = "\n\n".join(text_chunks)
            return DocumentParser.clean_extracted_text(combined)
        except Exception as exc:
            raise ValidationError(f"Unable to parse PDF content: {str(exc)}")

    @staticmethod
    def parse_docx(file_bytes: bytes) -> str:
        """Extract paragraphs and table text from DOCX bytes using python-docx."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_lines = []

            # Extract paragraphs
            for paragraph in doc.paragraphs:
                p_text = paragraph.text.strip()
                if p_text:
                    text_lines.append(p_text)

            # Extract table cells
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_lines.append(" | ".join(row_cells))

            combined = "\n".join(text_lines)
            return DocumentParser.clean_extracted_text(combined)
        except Exception as exc:
            raise ValidationError(f"Unable to parse DOCX content: {str(exc)}")

    @staticmethod
    def parse_txt(file_bytes: bytes) -> str:
        """Decode plain text bytes with fallback encodings."""
        for encoding in ["utf-8", "latin-1", "cp1252", "ascii"]:
            try:
                decoded = file_bytes.decode(encoding)
                return DocumentParser.clean_extracted_text(decoded)
            except UnicodeDecodeError:
                continue
        raise ValidationError("Unable to decode text file. Ensure it is encoded in UTF-8 or ASCII.")

    @staticmethod
    def clean_extracted_text(text: str) -> str:
        """Strip null bytes, non-printable characters, and normalize multiple line breaks."""
        text = text.replace("\x00", "")
        # Remove non-printable control characters except newline and tab
        text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", "", text)
        # Normalize whitespace and excessive newlines
        text = re.sub(r"\r\n|\r", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @classmethod
    def parse_document(cls, file_name: str, file_bytes: bytes) -> Tuple[str, str, int]:
        """Main entrypoint validating and parsing document into plain text."""
        file_type = cls.validate_file(file_name, file_bytes)

        if file_type == "pdf":
            raw_text = cls.parse_pdf(file_bytes)
        elif file_type == "docx":
            raw_text = cls.parse_docx(file_bytes)
        elif file_type == "txt":
            raw_text = cls.parse_txt(file_bytes)
        else:
            raise ValidationError(f"Unsupported file format: {file_type}")

        if not raw_text or len(raw_text.strip()) < 20:
            raise ValidationError(
                "Extracted text is empty or too short. If this is a scanned document, please ensure it has an OCR text layer."
            )

        return raw_text, file_type, len(file_bytes)
